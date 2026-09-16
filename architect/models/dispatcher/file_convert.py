"""Bounded shared-worker conversions; no agent execution or prose routing."""
from __future__ import annotations

import csv
import io
import re
import shutil
import subprocess
import tempfile
import uuid
import zipfile
from pathlib import Path
from typing import Any, Callable
from xml.etree import ElementTree

from fastapi import HTTPException
from pydantic import BaseModel, Field
import office_file

OFFICE = {"docx", "xlsx", "pptx"}
IMAGES = {"png", "jpg", "jpeg", "webp"}
INPUTS = OFFICE | IMAGES | {"pdf", "txt", "md", "csv"}
OUTPUTS = OFFICE | {"pdf", "png", "jpg", "txt", "csv"}
MAX_BYTES = 50 * 1024 * 1024
MAX_EXPANDED = 100 * 1024 * 1024
MAX_CELLS = 20000
MAX_PAGES = 100


class ConvertReq(BaseModel):
    source_path: str
    output_type: str
    mode: str = "auto"  # auto | content; content explicitly permits reflow
    pages: list[int] | None = Field(default=None, max_length=MAX_PAGES)
    dpi: int = Field(default=144, ge=72, le=300)
    send_zalo: bool = False
    thread_id: str = ""
    thread_type: str = "user"
    source_message_id: str = ""


def _check_office(path: Path) -> None:
    """Reject macros, external relationships, unsafe expansion before Office load."""
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > 10000 or sum(info.file_size for info in infos) > MAX_EXPANDED:
            raise ValueError("conversion_archive_limit")
        for info in infos:
            if info.filename.startswith("xl/externalLinks/") or info.filename == "xl/connections.xml":
                raise ValueError("conversion_external_resource_not_allowed")
            if info.filename.lower().endswith("vbaproject.bin"):
                raise ValueError("conversion_macros_not_allowed")
            if info.filename.endswith(".rels"):
                root = ElementTree.fromstring(archive.read(info))
                # Hyperlinks are inert metadata; externally fetched images,
                # workbook links and templates are not.
                for relation in root:
                    relation_type = relation.get("Type", "").lower()
                    if relation_type.endswith(("/vbaproject", "/oleobject", "/package")):
                        raise ValueError("conversion_active_content_not_allowed")
                    if relation.get("TargetMode") == "External" and not relation.get("Type", "").endswith("/hyperlink"):
                        raise ValueError("conversion_external_resource_not_allowed")
            if info.filename.startswith("xl/worksheets/") and info.filename.endswith(".xml"):
                root = ElementTree.fromstring(archive.read(info))
                for element in root.iter():
                    if element.tag.endswith("}f") and re.search(r"(?i)(?:^|[^A-Z0-9_])(?:_xlfn\.)?(?:WEBSERVICE|DDE|RTD)\s*\(", element.text or ""):
                        raise ValueError("conversion_external_resource_not_allowed")


def _to_pdf(source: Path, work: Path) -> Path:
    kind = source.suffix.lower().lstrip(".")
    dest = work / (source.stem + ".pdf")
    if kind == "pdf":
        shutil.copyfile(source, dest)
    elif kind in IMAGES:
        from PIL import Image
        with Image.open(source) as image:
            if image.width * image.height > 40_000_000:
                raise ValueError("conversion_image_limit")
            image.convert("RGB").save(dest, "PDF", resolution=144)
    elif kind in OFFICE:
        engine = shutil.which("libreoffice") or shutil.which("soffice")
        if not engine:
            raise ValueError("conversion_renderer_unavailable")
        # A per-request profile isolates concurrent conversions and prevents
        # connection to another worker's interactive LibreOffice process.
        subprocess.run([engine, "-env:UserInstallation=" + (work / "profile").as_uri(),
                        "--headless", "--convert-to", "pdf", "--outdir", str(work), str(source)],
                       check=True, timeout=120, capture_output=True)
    else:
        from html import escape
        text = source.read_text(encoding="utf-8-sig")
        office_file.write_pdf_from_html(dest, '<!DOCTYPE html><html><head><meta charset="utf-8"></head><body><pre style="white-space:pre-wrap">' + escape(text) + "</pre></body></html>")
    if not dest.is_file() or dest.stat().st_size == 0:
        raise ValueError("conversion_output_missing")
    return dest


def _content(source: Path) -> list[tuple[str, list[list[str]]]]:
    kind = source.suffix.lower().lstrip(".")
    if kind == "xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        try:
            count = sum(sheet.max_row * sheet.max_column for sheet in workbook)
            if count > MAX_CELLS:
                raise ValueError("conversion_cell_limit")
            return [(sheet.title, [["" if value is None else str(value) for value in row]
                                   for row in sheet.iter_rows(values_only=True)]) for sheet in workbook]
        finally:
            workbook.close()
    if kind == "docx":
        from docx import Document
        document = Document(source)
        # Walk the XML block order, not all paragraphs followed by all tables.
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        rows = []
        for block in document.element.body:
            if block.tag.endswith("}p"):
                rows.append([Paragraph(block, document).text])
            elif block.tag.endswith("}tbl"):
                rows.extend([[cell.text for cell in row.cells] for row in Table(block, document).rows])
        return [(source.stem, rows)]
    if kind == "pptx":
        from pptx import Presentation
        deck = Presentation(source)
        if len(deck.slides) > MAX_PAGES:
            raise ValueError("conversion_page_limit")
        sections = []
        for index, slide in enumerate(deck.slides, 1):
            rows = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    rows.append([shape.text])
                elif shape.has_table:
                    rows.extend([[cell.text for cell in row.cells] for row in shape.table.rows])
            sections.append((f"Slide {index}", rows))
        return sections
    if kind == "pdf":
        import pymupdf
        with pymupdf.open(source) as document:
            if len(document) > MAX_PAGES or document.needs_pass:
                raise ValueError("conversion_pdf_limit_or_password")
            sections = []
            for index, page in enumerate(document, 1):
                text = page.get_text().strip()
                if not text:
                    raise ValueError("conversion_ocr_required")
                sections.append((f"Page {index}", [[line] for line in text.splitlines()]))
            return sections
    if kind in {"txt", "md", "csv"}:
        text = source.read_text(encoding="utf-8-sig")
        rows = list(csv.reader(io.StringIO(text))) if kind == "csv" else [[line] for line in text.splitlines()]
        return [(source.stem, rows)]
    raise ValueError("conversion_ocr_required")


def _write_content(dest: Path, sections: list[tuple[str, list[list[str]]]]) -> None:
    kind = dest.suffix.lstrip(".")
    if sum(len(row) for _, rows in sections for row in rows) > MAX_CELLS:
        raise ValueError("conversion_cell_limit")
    if kind == "docx":
        from docx import Document
        document = Document()
        for title, rows in sections:
            document.add_heading(title, level=1)
            for row in rows:
                if len(row) == 1:
                    document.add_paragraph(row[0])
                else:
                    cells = document.add_table(rows=1, cols=len(row)).rows[0].cells
                    for cell, value in zip(cells, row):
                        cell.text = value
        document.save(dest)
    elif kind == "xlsx":
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.remove(workbook.active)
        for index, (title, rows) in enumerate(sections, 1):
            # Stable numeric prefix also avoids duplicate worksheet names.
            name = str(index) + " " + "".join(c for c in title if c not in "[]:*?/\\")
            sheet = workbook.create_sheet(name[:31])
            for row in rows:
                sheet.append(row)
                # Imported strings remain data, not injected formulas. Source
                # workbook formulas are preserved as visible formula text.
                for cell in sheet[sheet.max_row]:
                    cell.data_type = "s"
        workbook.save(dest)
    elif kind == "pptx":
        from pptx import Presentation
        from pptx.util import Inches, Pt
        deck = Presentation()
        for title, rows in sections:
            lines = [" | ".join(row) for row in rows]
            if any(len(line) > 160 for line in lines):
                raise ValueError("conversion_presentation_reflow_required")
            chunks = [lines[i:i + 8] for i in range(0, len(lines), 8)] or [[]]
            for chunk in chunks:
                if len(deck.slides) >= MAX_PAGES:
                    raise ValueError("conversion_page_limit")
                slide = deck.slides.add_slide(deck.slide_layouts[5])
                slide.shapes.title.text = title
                frame = slide.shapes.add_textbox(Inches(.6), Inches(1.3), Inches(8.8), Inches(5.5)).text_frame
                frame.word_wrap = True
                for index, line in enumerate(chunk):
                    paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                    paragraph.text = line
                    paragraph.font.size = Pt(18)
        deck.save(dest)
    elif kind == "csv":
        if len(sections) != 1:
            raise ValueError("conversion_csv_requires_single_section")
        with dest.open("w", encoding="utf-8-sig", newline="") as output:
            csv.writer(output).writerows(sections[0][1])
    else:
        dest.write_text("\n\n".join(title + "\n" + "\n".join("\t".join(row) for row in rows)
                                    for title, rows in sections), encoding="utf-8")


def convert_file(req: ConvertReq, work: Path) -> tuple[list[Path], str]:
    source = office_file._resolve_media_path(req.source_path)
    kind = req.output_type.lower().lstrip(".")
    if source is None or not source.is_file() or source.suffix.lower().lstrip(".") not in INPUTS:
        raise ValueError("conversion_source_unavailable_or_unsupported")
    if source.stat().st_size > MAX_BYTES:
        raise ValueError("conversion_source_limit")
    if kind not in OUTPUTS or req.mode not in {"auto", "content"}:
        raise ValueError("conversion_target_or_mode_unsupported")
    source_kind = source.suffix.lower().lstrip(".")
    if source_kind in OFFICE:
        _check_office(source)
    if req.pages is not None and kind not in {"png", "jpg"}:
        raise ValueError("conversion_pages_only_for_images")
    if kind == source_kind and req.pages is None:
        dest = work / ("converted." + kind)
        shutil.copyfile(source, dest)
        return [dest], "copy"
    if kind == "pdf" or kind in {"png", "jpg"}:
        pdf = _to_pdf(source, work)
        import pymupdf
        with pymupdf.open(pdf) as document:
            if document.needs_pass or len(document) > MAX_PAGES:
                raise ValueError("conversion_pdf_limit_or_password")
            if kind == "pdf":
                return [pdf], "print-layout"
            pages = req.pages if req.pages is not None else list(range(1, len(document) + 1))
            if not pages or len(set(pages)) != len(pages) or any(page < 1 or page > len(document) for page in pages):
                raise ValueError("conversion_page_selection_invalid")
            outputs = []
            pixels = 0
            for number in pages:
                page = document[number - 1]
                pixels += page.rect.width * page.rect.height * (req.dpi / 72) ** 2
                if pixels > 80_000_000:
                    raise ValueError("conversion_raster_limit")
                dest = work / f"page-{number:03d}.{kind}"
                page.get_pixmap(dpi=req.dpi, alpha=False).save(dest)
                outputs.append(dest)
            return outputs, "page-raster"
    if req.mode != "content":
        raise ValueError("conversion_content_mode_required")
    dest = work / ("converted." + kind)
    _write_content(dest, _content(source))
    return [dest], "content-reflow"


def register_file_convert(app: Any, media_dir: Path, deliver: Callable[..., dict],
                          scan: Callable[[Path, str, str], str] | None = None) -> None:
    @app.post("/v1/file-convert")
    def endpoint(req: ConvertReq) -> dict:
        if not office_file._enabled():
            raise HTTPException(503, "File conversion is unavailable.")
        if req.send_zalo and not req.thread_id:
            raise HTTPException(400, "thread_id required")
        if req.send_zalo and (not req.source_message_id.strip() or req.thread_type not in {"user", "group"}):
            raise HTTPException(400, "original messaging delivery binding required")
        source = office_file._resolve_media_path(req.source_path)
        if source is None or not source.is_file() or source.suffix.lower().lstrip(".") not in INPUTS:
            raise HTTPException(422, "conversion_source_unavailable_or_unsupported")
        if source.stat().st_size > MAX_BYTES:
            raise HTTPException(422, "conversion_source_limit")
        # Reuse the deployed scan policy before parsing untrusted documents,
        # including direct sessions without a messaging destination.
        if scan is not None and scan(source, req.thread_id, source.name) == "blocked":
            raise HTTPException(403, "File conversion was blocked by the security policy.")
        build = media_dir / "out/.build"
        files = []
        try:
            build.mkdir(parents=True, exist_ok=True)
            with tempfile.TemporaryDirectory(dir=build) as temp:
                outputs, fidelity = convert_file(req, Path(temp))
                # Publish a whole successful conversion together, outside the
                # public watcher. Native clients or explicit host send own delivery.
                publication = media_dir / "out/.conversions" / uuid.uuid4().hex
                publication.parent.mkdir(parents=True, exist_ok=True)
                staged = Path(temp) / "publication"
                staged.mkdir()
                for output in outputs:
                    dest = publication / output.name
                    output.replace(staged / output.name)
                    files.append({"path": str(dest), "hermes_path": str(Path("/opt/data/media") / dest.relative_to(media_dir)), "file": dest.name})
                # A single same-filesystem directory rename exposes the complete
                # batch. A failed second page never leaves a partial publication.
                staged.replace(publication)
        except (ValueError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
            raise HTTPException(422, str(exc)) from exc
        except subprocess.TimeoutExpired as exc:
            raise HTTPException(504, "conversion_timeout") from exc
        except subprocess.CalledProcessError as exc:
            raise HTTPException(422, "conversion_engine_failed") from exc
        except OSError as exc:
            raise HTTPException(503, "conversion_storage_or_engine_unavailable") from exc
        for item in files:
            if req.send_zalo:
                item["delivery"] = deliver(path=item["path"], thread_id=req.thread_id,
                                           thread_type=req.thread_type,
                                           source_message_id=req.source_message_id.strip(),
                                           caption="", filename=item["file"], lock_thread=True)
        return {"ok": True, "fidelity": fidelity, "files": files,
                "warnings": ["Editable content was reflowed; original layout, images, charts, animations and formulas are not preserved."] if fidelity == "content-reflow" else []}
