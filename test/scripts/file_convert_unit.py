"""Real conversion engines, content integrity and destination-free HTTP boundary."""
import os
from pathlib import Path
import sys
import tempfile
from unittest.mock import Mock, patch
import zipfile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "architect/models/dispatcher"))
import file_convert
import office_file


def main():
    import pymupdf
    from docx import Document
    from openpyxl import Workbook, load_workbook
    from pptx import Presentation
    from PIL import Image
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    with tempfile.TemporaryDirectory() as temp:
        root = Path(temp)
        document = Document()
        document.add_paragraph("Thông tin tiếng Việt")
        document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Sample value"
        document.save(root / "sample.docx")
        workbook = Workbook()
        workbook.active.title = "Alpha"
        workbook.active.append(["Sample value", "=1+2"])
        workbook.create_sheet("Beta").append(["Other content", 123])
        for sheet in workbook:
            sheet.column_dimensions["A"].width = 24
            sheet.column_dimensions["B"].width = 16
        workbook.save(root / "sample.xlsx")
        deck = Presentation()
        for text in ("First slide", "Second slide"):
            deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = text
        deck.save(root / "sample.pptx")
        Image.new("RGB", (320, 180), "blue").save(root / "sample.png")
        count = 0
        with patch.object(office_file, "_MEDIA_ROOTS", (root,)), patch.dict(os.environ, {"MEDIA_CACHE_DIR": str(root), "OFFICE_FILE_GEN": "active"}):
            def run(name):
                nonlocal count
                count += 1
                print(f"running test case {count}/12 {name}", flush=True)
            for kind in ("docx", "xlsx", "pptx"):
                run(kind + " to PDF with the actual print engine")
                work = root / kind
                work.mkdir()
                outputs, fidelity = file_convert.convert_file(file_convert.ConvertReq(source_path=str(root / ("sample." + kind)), output_type="pdf"), work)
                assert fidelity == "print-layout"
                with pymupdf.open(outputs[0]) as pdf:
                    text = "".join(page.get_text() for page in pdf)
                    if kind == "pptx":
                        assert len(pdf) == 2 and "Second slide" in pdf[1].get_text()
                    else:
                        assert "Sample value" in text
                        if kind == "xlsx":
                            assert "Other content" in text
            run("PPTX to selected image, no intermediate PDF delivered")
            work = root / "raster"
            work.mkdir()
            outputs, _ = file_convert.convert_file(file_convert.ConvertReq(source_path=str(root / "sample.pptx"), output_type="png", pages=[2]), work)
            assert len(outputs) == 1 and outputs[0].name == "page-002.png"
            with Image.open(outputs[0]) as image:
                assert image.width > 1000
            run("PPTX to all PNG slides, ordered and visually distinct")
            outputs, fidelity = file_convert.convert_file(file_convert.ConvertReq(source_path=str(root / "sample.pptx"), output_type="png"), work)
            assert len(outputs) == 2 and [p.name for p in outputs] == ["page-001.png", "page-002.png"]
            assert outputs[0].read_bytes() != outputs[1].read_bytes()
            for path in outputs:
                with Image.open(path) as image:
                    assert image.width > 1000 and image.height > 500
            run("PDF to all JPG pages and selected PNG page")
            pdf_source = root / "pptx" / "sample.pdf"
            all_images, _ = file_convert.convert_file(file_convert.ConvertReq(source_path=str(pdf_source), output_type="jpg"), work)
            assert len(all_images) == 2 and [p.name for p in all_images] == ["page-001.jpg", "page-002.jpg"]
            selected, _ = file_convert.convert_file(file_convert.ConvertReq(source_path=str(pdf_source), output_type="png", pages=[2]), work)
            assert len(selected) == 1 and selected[0].name == "page-002.png"
            with pymupdf.open(pdf_source) as original, Image.open(selected[0]) as image:
                assert abs(image.width / image.height - original[1].rect.width / original[1].rect.height) < 0.01
            run("Excel to Word preserves both sheets and formula text")
            work = root / "reflow"
            work.mkdir()
            outputs, _ = file_convert.convert_file(file_convert.ConvertReq(source_path=str(root / "sample.xlsx"), output_type="docx", mode="content"), work)
            doc = Document(outputs[0])
            text = " ".join(p.text for p in doc.paragraphs) + " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            assert all(value in text for value in ("Alpha", "Beta", "=1+2", "Other content"))
            run("Word to Excel keeps imported formula-like strings inert")
            sample = Document()
            sample.add_paragraph("=HYPERLINK(\"https://example.com\")")
            sample.save(root / "formula.docx")
            outputs, _ = file_convert.convert_file(file_convert.ConvertReq(source_path=str(root / "formula.docx"), output_type="xlsx", mode="content"), work)
            book = load_workbook(outputs[0])
            assert book.active["A1"].data_type == "s"
            run("image to PDF retains an actual image")
            outputs, _ = file_convert.convert_file(file_convert.ConvertReq(source_path=str(root / "sample.png"), output_type="pdf"), work)
            with pymupdf.open(outputs[0]) as pdf:
                assert pdf[0].get_image_info()
            run("lossy, invalid page and external paths fail closed")
            for req in (file_convert.ConvertReq(source_path=str(root / "sample.xlsx"), output_type="docx"),
                        file_convert.ConvertReq(source_path=str(root / "sample.pptx"), output_type="png", pages=[0]),
                        file_convert.ConvertReq(source_path="/etc/passwd", output_type="txt", mode="content")):
                try:
                    file_convert.convert_file(req, work)
                    raise AssertionError("unsafe conversion accepted")
                except ValueError:
                    pass
            run("Office external resources rejected before renderer execution")
            unsafe = root / "external.docx"
            with zipfile.ZipFile(unsafe, "w") as archive:
                archive.writestr("word/_rels/document.xml.rels", '<Relationships><Relationship TargetMode="External" Type="sample/image" Target="http://127.0.0.1/private"/></Relationships>')
            try:
                file_convert.convert_file(file_convert.ConvertReq(source_path=str(unsafe), output_type="pdf"), work)
                raise AssertionError("external resource accepted")
            except ValueError as exc:
                assert str(exc) == "conversion_external_resource_not_allowed"
            run("direct API has no messaging dependency and outputs stay private")
            app = FastAPI()
            delivery = Mock(side_effect=AssertionError("unexpected Zalo dependency"))
            file_convert.register_file_convert(app, root, delivery)
            response = TestClient(app).post("/v1/file-convert", json={"source_path": str(root / "sample.png"), "output_type": "pdf"})
            assert response.status_code == 200, response.text
            result = response.json()
            assert len(result["files"]) == 1 and "/.conversions/" in result["files"][0]["hermes_path"]
            assert Path(result["files"][0]["path"]).is_file()
            delivery.assert_not_called()
            blocked_app = FastAPI()
            scan = Mock(return_value="blocked")
            file_convert.register_file_convert(blocked_app, root, delivery, scan)
            with patch.object(file_convert, "convert_file", side_effect=AssertionError("blocked source parsed")):
                blocked = TestClient(blocked_app).post("/v1/file-convert", json={"source_path": str(root / "sample.docx"), "output_type": "pdf"})
                assert blocked.status_code == 403
                scan.assert_called_once()
                delivery.assert_not_called()
    print("PASS real file conversions")


if __name__ == "__main__":
    main()
