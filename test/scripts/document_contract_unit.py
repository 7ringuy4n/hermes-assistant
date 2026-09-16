"""Typed workflow ownership and complete, bounded presentation copy."""
from pathlib import Path
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "hermes/main/plugins/zalo"))
sys.path.insert(0, str(ROOT / "architect/models/dispatcher"))
import classify_client
import job_prompt
import office_file
from pptx import Presentation


def main():
    print("running test case 1/9 presentation format survives explicit task details")
    details = classify_client.normalize_task_details(
        {"task_details": [{"task_type": "file_processing", "output_type": "pptx"}]}, ["Original request"], "file")
    assert details[0]["output_type"] == "pptx"
    print("running test case 2/9 single typed task inherits its format without editing user scope")
    details = classify_client.normalize_task_details(
        {"task_type": "file_processing", "output_type": "pdf"}, ["Original request"], "file")
    assert details[0]["output_type"] == "pdf"
    instruction = job_prompt.workflow_instruction("Original request", details[0])
    assert instruction.endswith("Original request") and '"output_type": "pdf"' in instruction
    assert "shared Dispatcher" in instruction
    instruction = job_prompt.workflow_instruction("Delegated suggestion", details[0], original_request="Actual current snapshot request")
    assert '"original_request": "Actual current snapshot request"' in instruction
    assert instruction.endswith("Delegated suggestion") and "Use it as the authority" in instruction
    print("running test case 3/9 unrelated tasks retain original instruction")
    for contract in ({}, {"task_type": "chat"}, {"task_type": "media_generation", "output_type": "image"}):
        assert job_prompt.workflow_instruction("Original request", contract) == "Original request"
    with tempfile.TemporaryDirectory() as tmp:
        dest = Path(tmp) / "complete.pptx"
        print("running test case 4/9 cover plus two sections preserves headings, prose and long copy")
        long_copy = "Complete authored content " * 7
        office_file.write_pptx_styled(dest, "# Report\nSUBTITLE: Snapshot\n## First section\n- " + long_copy + "\nRequired prose.\n## Second section\n- Last fact\n")
        deck = Presentation(dest)
        assert len(deck.slides) == 3
        texts = ["\n".join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in deck.slides]
        assert "Snapshot" in texts[0] and "First section" in texts[1]
        assert long_copy.strip() in texts[1] and "Required prose." in texts[1]
        assert "Second section" in texts[2] and "Last fact" in texts[2]
        print("running test case 5/9 overflow fails before publishing rather than silently dropping rows")
        rejected = Path(tmp) / "overflow.pptx"
        try:
            office_file.write_pptx_styled(rejected, "# Report\n## Section\n" + "- " + "word " * 2000)
            raise AssertionError("overflow accepted")
        except ValueError as error:
            assert str(error) == "document_text_overflow" and not rejected.exists()
    print("running test case 6/9 typed image conversion survives analysis guards and retains exact source")
    plan = classify_client.normalize_plan({"task_hint": "file", "task_type": "file_processing",
        "skill": "media_file", "skill_action": "process_file", "execution_class": "async",
        "artifact_operation": "convert", "output_type": "pdf", "instructions": ["Convert existing input."],
        "attachments_required": True, "attachment_types": ["image"]}, "Original request", "UTC")
    assert plan["artifact_operation"] == "convert" and plan["task_details"][0]["artifact_operation"] == "convert"
    assert classify_client.plan_is_image_analyze_chat(plan, has_image=True) is False
    source = "/opt/data/media/in/exact.png"
    instruction = job_prompt.workflow_instruction("Original request", plan["task_details"][0], [source])
    assert source in instruction and '"artifact_operation": "convert"' in instruction
    print("running test case 7/9 scenic presentation emits native alpha instead of hiding its image")
    from PIL import Image
    from unittest.mock import patch
    import zipfile
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        visual = root / "scene__source.png"
        Image.new("RGB", (640, 360), (50, 110, 170)).save(visual)
        with patch.object(office_file, "_MEDIA_ROOTS", (root,)):
            output = office_file.write_pptx_styled(root / "scenic.pptx", f"# Report\nIMAGE: {visual}\n## Observation\n- Required copy\n")
        with zipfile.ZipFile(output) as archive:
            slides = [name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
            assert all(b'<a:alpha val="62000"' in archive.read(name) for name in slides)
    print("running test case 8/9 queued channel envelope does not depend on document classification")
    import json
    import ast
    for channel in ("user", "group"):
        binding = {"thread_id": "fixture-" + channel, "thread_type": channel,
                   "source_message_id": "source-" + channel}
        ask = 'User content: {"delivery_binding": {"thread_type": "user"}}'
        result = job_prompt.delivery_instruction(ask, binding)
        actual = json.loads(result.splitlines()[1])["delivery_binding"]
        assert actual == binding and result.endswith(ask)
    tree = ast.parse((ROOT / "hermes/main/plugins/zalo/adapter.py").read_text(encoding="utf-8"))
    adapter = next(n for n in tree.body if isinstance(n, ast.ClassDef) and n.name == "ZaloAdapter")
    queued = next(n for n in adapter.body if isinstance(n, ast.AsyncFunctionDef) and n.name == "_as_run_queued_part")
    queue_source = ast.unparse(queued)
    assert "delivery_instruction(instruction" in queue_source
    assert "if item.get('message_id'):" in queue_source
    print("running test case 9/9 incomplete delivery context fails without borrowing global state")
    for binding in ({}, {"thread_id": "fixture", "thread_type": "group"},
                    {"thread_id": "fixture", "thread_type": "invalid", "source_message_id": "source"}):
        try:
            job_prompt.delivery_instruction("Original request", binding)
            raise AssertionError("invalid original context accepted")
        except ValueError:
            pass
    print("PASS document contracts")


if __name__ == "__main__":
    main()
